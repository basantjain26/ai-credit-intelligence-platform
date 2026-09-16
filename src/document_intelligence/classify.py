from pathlib import Path
import json
import csv
from zipfile import ZipFile
import fitz
import psycopg2
from docx import Document
from openpyxl import load_workbook


DB_CONFIG = {
    "host": "127.0.0.1",
    "port": 5434,
    "dbname": "credit_intelligence",
    "user": "credit_user",
    "password": "credit_password",
}

DOCUMENT_TYPE_RULES = {
    "FINANCIAL_STATEMENT": {
        "folder_keywords": ["financial_statements"],
        "filename_keywords": [
            "financial",
            "audited",
            "balance_sheet",
            "income_statement",
        ],
        "content_keywords": [
            "revenue",
            "ebitda",
            "net income",
            "balance sheet",
            "cash flow",
        ],
    },
    "BORROWER_PROFILE": {
        "folder_keywords": ["borrower_profiles"],
        "filename_keywords": [
            "borrower",
            "profile",
            "credit_note",
        ],
        "content_keywords": [
            "borrower",
            "company overview",
            "credit request",
            "management",
        ],
    },
    "EXPOSURE_SCHEDULE": {
        "folder_keywords": ["exposure_schedules"],
        "filename_keywords": [
            "exposure",
            "facility",
            "loan_schedule",
        ],
        "content_keywords": [
            "outstanding balance",
            "interest rate",
            "maturity",
            "secured",
        ],
    },
    "TRANSACTION_EXTRACT": {
        "folder_keywords": ["transaction_extracts"],
        "filename_keywords": [
            "transaction",
            "bank_activity",
        ],
        "content_keywords": [
            "transaction_id",
            "amount",
            "counterparty",
            "direction",
        ],
    },
    "CREDIT_BUREAU_REPORT": {
        "folder_keywords": ["credit_bureau"],
        "filename_keywords": [
            "bureau",
            "credit_report",
        ],
        "content_keywords": [
            "credit score",
            "credit_score",
            "delinquencies",
            "credit exposure",
        ],
    },
    "LENDING_POLICY": {
        "folder_keywords": ["policies"],
        "filename_keywords": [
            "policy",
            "lending_policy",
            "commercial_lending",
        ],
        "content_keywords": [
            "commercial lending policy",
            "debt service coverage",
            "dscr",
            "policy exception",
            "related party",
            "credit officer",
        ],
    },
}

def normalize(text: str) -> str:
    return text.lower().replace("-", " ").replace("_", " ")


def score_keywords(text: str, keywords: list[str], weight: int) -> int:
    text = normalize(text)

    score = 0

    for keyword in keywords:
        if normalize(keyword) in text:
            score += weight

    return score

def read_pdf_preview(file_path: Path) -> str:
    document = fitz.open(file_path)

    text_parts = []

    for page_number in range(min(2, len(document))):
        page = document[page_number]
        text_parts.append(page.get_text())

    document.close()

    return "\n".join(text_parts)[:5000]

def read_docx_preview(file_path: Path) -> str:
    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs[:30]:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    return "\n".join(paragraphs)[:5000]

def read_xlsx_preview(file_path: Path) -> str:
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    values = []

    for sheet in workbook.worksheets[:2]:

        values.append(sheet.title)

        for row_index, row in enumerate(
            sheet.iter_rows(values_only=True)
        ):
            row_values = [
                str(value)
                for value in row
                if value is not None
            ]

            values.append(" ".join(row_values))

            if row_index >= 15:
                break

    workbook.close()

    return "\n".join(values)[:5000]

def read_text_preview(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    try:
        if suffix == ".json":
            with open(file_path, "r", encoding="utf-8") as file:
                data = json.load(file)

            return json.dumps(data)[:5000]

        if suffix == ".csv":
            lines = []

            with open(file_path, "r", encoding="utf-8") as file:
                reader = csv.reader(file)

                for index, row in enumerate(reader):
                    lines.append(" ".join(row))

                    if index >= 10:
                        break

            return " ".join(lines)
        
        if suffix == ".txt":
            return file_path.read_text(
                encoding="utf-8")[:5000]
        
        if suffix == ".docx":
            return read_docx_preview(file_path)

        if suffix == ".xlsx":
            return read_xlsx_preview(file_path)

        if suffix == ".pdf":
            return read_pdf_preview(file_path)

    except Exception as exc:
        print(f"Preview extraction failed for {file_path.name}: {exc}")

    return ""




def classify_document(file_path: Path) -> dict:
    folder_name = file_path.parent.name
    filename = file_path.name
    content = read_text_preview(file_path)

    scores = {}

    for document_type, rules in DOCUMENT_TYPE_RULES.items():

        score = 0

        score += score_keywords(
            folder_name,
            rules["folder_keywords"],
            weight=5,
        )

        score += score_keywords(
            filename,
            rules["filename_keywords"],
            weight=3,
        )

        score += score_keywords(
            content,
            rules["content_keywords"],
            weight=1,
        )

        scores[document_type] = score

    best_type = max(scores, key=scores.get)
    best_score = scores[best_type]

    sorted_scores = sorted(
        scores.values(),
        reverse=True,
    )

    second_best_score = (
        sorted_scores[1]
        if len(sorted_scores) > 1
        else 0
    )

    if best_score == 0:
        return {
            "document_type": "UNKNOWN",
            "confidence": 0.0,
            "method": "DETERMINISTIC_RULES",
            "scores": scores,
        }

    margin = best_score - second_best_score

    confidence = min(
        1.0,
        (best_score / 10) + (margin / 20),
    )

    return {
        "document_type": best_type,
        "confidence": round(confidence, 3),
        "method": "DETERMINISTIC_RULES",
        "scores": scores,
    }


SOURCE_ROOT = Path("data/source/documents")

def save_classification(
    document_name: str,
    result: dict,
) -> None:

    connection = psycopg2.connect(**DB_CONFIG)

    try:
        with connection:
            with connection.cursor() as cursor:

                cursor.execute(
                    """
                    UPDATE documents
                    SET
                        document_type = %s,
                        classification_method = %s,
                        classification_confidence = %s,
                        classified_at = CURRENT_TIMESTAMP,
                        processing_status = 'CLASSIFIED'
                    WHERE document_name = %s
                    """,
                    (
                        result["document_type"],
                        result["method"],
                        result["confidence"],
                        document_name,
                    ),
                )

    finally:
        connection.close()


def main():
    for file_path in SOURCE_ROOT.rglob("*"):

        if not file_path.is_file():
            continue

        result = classify_document(file_path)

        save_classification(
            document_name=file_path.name,
            result=result,
        )

        print()
        print(file_path.name)
        print(f"Type       : {result['document_type']}")
        print(f"Confidence : {result['confidence']}")
        print(f"Method     : {result['method']}")
        print(f"Scores     : {result['scores']}")


if __name__ == "__main__":
    main()



