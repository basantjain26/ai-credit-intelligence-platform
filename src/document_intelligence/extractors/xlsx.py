from pathlib import Path
import uuid

from openpyxl import load_workbook

from .base import DocumentExtractor
from src.document_intelligence.models import ContentBlock, ExtractedDocument


class XLSXExtractor(DocumentExtractor):

    def extract(
        self,
        file_path: Path,
        document_id: str,
    ) -> ExtractedDocument:

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        content_blocks = []

        try:
            for worksheet in workbook.worksheets:

                rows = []

                for row in worksheet.iter_rows(values_only=True):
                    rows.append(
                        [
                            value
                            for value in row
                        ]
                    )

                content_blocks.append(
                    ContentBlock(
                        block_id=self._new_block_id(document_id),
                        block_type="TABLE",
                        rows=rows,
                        sheet_name=worksheet.title,
                        table_index=1,
                    )
                )

        finally:
            workbook.close()

        return ExtractedDocument(
            document_id=document_id,
            file_name=file_path.name,
            file_type="XLSX",
            parser_name="openpyxl",
            parser_version="1.x",
            content_blocks=content_blocks,
        )

    @staticmethod
    def _new_block_id(document_id: str) -> str:
        suffix = uuid.uuid4().hex[:10].upper()
        return f"{document_id}_BLK_{suffix}"