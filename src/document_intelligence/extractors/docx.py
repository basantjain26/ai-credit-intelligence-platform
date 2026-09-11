from pathlib import Path
import uuid
from docx import Document

from .base import DocumentExtractor
from src.document_intelligence.models import ContentBlock, ExtractedDocument


class DOCXExtractor(DocumentExtractor):

    def extract(
        self,
        file_path: Path,
        document_id: str,
    ) -> ExtractedDocument:

        document = Document(file_path)
        content_blocks = []

        for paragraph_index, paragraph in enumerate(
            document.paragraphs
        ):
            text = paragraph.text.strip()

            if not text:
                continue

            content_blocks.append(
                ContentBlock(
                    block_id=self._new_block_id(document_id),
                    block_type="TEXT",
                    text=text,
                    paragraph_index=paragraph_index,
                )
            )

        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):
            rows = []

            for row in table.rows:
                rows.append(
                    [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                )

            content_blocks.append(
                ContentBlock(
                    block_id=self._new_block_id(document_id),
                    block_type="TABLE",
                    rows=rows,
                    table_index=table_index,
                )
            )

        return ExtractedDocument(
            document_id=document_id,
            file_name=file_path.name,
            file_type="DOCX",
            parser_name="python-docx",
            parser_version="1.x",
            content_blocks=content_blocks,
        )

    @staticmethod
    def _new_block_id(document_id: str) -> str:
        suffix = uuid.uuid4().hex[:10].upper()
        return f"{document_id}_BLK_{suffix}"